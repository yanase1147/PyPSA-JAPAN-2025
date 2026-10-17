"""FIT/FIP の公表統計から、モデルに入っていない中小のバイオマス発電をエリア別に集計し、
入力Excel にエリアごとの発電機「Biomass (small, <エリア>)」として追加する。

使い方:
    python tools/small_biomass.py data/Network_config_2025.xlsx A_pref202512.xlsx C_kaitori202603.csv [--year 2025]

入力:
- A表（都道府県別認定・導入量, xlsx）: 「表A②－１」新規認定分・「表A②－２」移行認定分の導入容量のうち、
  「バイオマス比率考慮あり」の列（石炭混焼はバイオマス相当分の容量だけを数える）
- C表（買取電力量, csv）: バイオマス発電の月別買取電力量（万kWh）

考え方:
1. 都道府県の導入容量（比率考慮あり、新規＋移行）を9エリアに集計する（沖縄は除く）。
2. 対象年の買取電力量（全国、暦年）を、エリアの容量の比率で配分する。
3. モデルにすでにある大型バイオマス（generators シートの carrier=Biomass）の容量と発電量を差し引き、
   残りを中小のバイオマスとしてエリアごとに1つの発電機にする。
4. 出力は C表の月別の発電量の形に沿って、月内は一定で固定する（biomass_cf シートに列を追加。
   build_network はバイオマスを biomass_cf の値に固定する）。

注意:
- 都道府県のエリアへの対応は簡易化している（静岡は中部、福井は北陸、三重は中部など、県内の主なエリアに寄せる）。
- C表は FIT の買取分で、FIP に移行した設備の発電量は含まない。
- 石炭混焼のバイオマス相当分も含む。混焼の発電所はモデルでは石炭として扱っているため、その分は二重に数えている。
"""
import argparse
import csv

import numpy as np
import openpyxl
import pandas as pd

AREA = {
    "北海道": "Hokkaido",
    **{p: "Tohoku" for p in ["青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県", "新潟県"]},
    **{p: "Tokyo" for p in ["茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県", "山梨県"]},
    **{p: "Chubu" for p in ["長野県", "岐阜県", "静岡県", "愛知県", "三重県"]},
    **{p: "Hokuriku" for p in ["富山県", "石川県", "福井県"]},
    **{p: "Kansai" for p in ["滋賀県", "京都府", "大阪府", "兵庫県", "奈良県", "和歌山県"]},
    **{p: "Chugoku" for p in ["鳥取県", "島根県", "岡山県", "広島県", "山口県"]},
    **{p: "Shikoku" for p in ["徳島県", "香川県", "愛媛県", "高知県"]},
    **{p: "Kyusyu" for p in ["福岡県", "佐賀県", "長崎県", "熊本県", "大分県", "宮崎県", "鹿児島県"]},
    "沖縄県": None,
}
MARKER = "バイオマス比率考慮あり"
REIWA = {2024: "令和6年", 2025: "令和7年", 2026: "令和8年"}


def capacity_by_pref(a_xlsx):
    """表A②－１・②－２ の「バイオマス比率考慮あり」6列の合計 [MW] を都道府県別に返す。"""
    wb = openpyxl.load_workbook(a_xlsx, data_only=True)
    total = pd.Series(0.0, index=list(AREA))
    for name in ["表A②－１", "表A②－２"]:
        ws = wb[name]
        c0 = next(c for c in range(1, ws.max_column + 1) if MARKER in str(ws.cell(3, c).value or ""))
        for r in range(6, ws.max_row + 1):
            pref = ws.cell(r, 1).value
            if pref in AREA:
                total[pref] += sum(float(ws.cell(r, c).value or 0) for c in range(c0, c0 + 6)) / 1e3  # kW→MW
    return total


def monthly_energy(c_csv, year):
    """C表のバイオマス発電の対象年1〜12月の買取電力量 [MWh] を返す。"""
    rows = list(csv.reader(open(c_csv, encoding="utf-8-sig")))
    yrow = next(r for r in rows if sum("年度" in c for c in r) > 5)
    mrow = next(r for r in rows if sum("月分" in c for c in r) > 10)
    bio = next(r for r in rows if len(r) > 1 and r[1] == "バイオマス発電")
    vals = [float(b.replace(",", "")) * 10 for y, m, b in zip(yrow, mrow, bio)  # 万kWh→MWh
            if y.startswith(REIWA[year]) and "年度" not in y and m]
    if len(vals) != 12:
        raise ValueError(f"C表に {year} 年の12か月分がありません（{len(vals)}か月）")
    return pd.Series(vals, index=range(1, 13))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("a_xlsx")
    ap.add_argument("c_csv")
    ap.add_argument("--year", type=int, default=2025)
    args = ap.parse_args()

    cap_pref = capacity_by_pref(args.a_xlsx)
    month = monthly_energy(args.c_csv, args.year)
    cap_area = cap_pref.groupby(cap_pref.index.map(AREA)).sum()            # 沖縄は index が None で落ちる
    japan_ex_okinawa = cap_area.sum() / cap_pref.sum()
    energy_year = month.sum() * japan_ex_okinawa                           # 沖縄分を容量比で除く [MWh]
    energy_area = energy_year * cap_area / cap_area.sum()

    x = pd.read_excel(args.xlsx, sheet_name=["generators", "biomass_cf"])
    g, bcf = x["generators"], x["biomass_cf"]
    big = g[(g.carrier == "Biomass") & ~g.name.str.startswith("Biomass (small")].set_index("name")
    snaps = pd.date_range(f"{args.year}-01-01", periods=8760, freq="h")
    avg = (bcf[bcf.columns.intersection(big.index)] * big.p_nom).sum(axis=1) / big.p_nom[bcf.columns.intersection(big.index)].sum()
    big_e = pd.Series({n: (bcf[n] if n in bcf else avg).sum() * big.p_nom[n] for n in big.index})
    big_cap_area = big.p_nom.groupby(big.area).sum().reindex(cap_area.index).fillna(0)
    big_e_area = big_e.groupby(big.area).sum().reindex(cap_area.index).fillna(0)

    net_cap = (cap_area - big_cap_area).clip(lower=0)
    net_e = (energy_area - big_e_area).clip(lower=0)
    shape = month / month.mean()                                           # 月別の形（平均1）
    hourly_shape = shape.reindex(snaps.month).values

    new_rows, new_cols, report = [], {}, []
    for area in cap_area.index:
        if net_cap[area] <= 0 or net_e[area] <= 0:
            continue
        name = f"Biomass (small, {area})"
        cf_mean = min(net_e[area] / (net_cap[area] * 8760), 0.95)
        prof = np.clip(hourly_shape * cf_mean, 0, 1)
        new_cols[name] = np.round(prof, 6)
        new_rows.append({"name": name, "area": area, "carrier": "Biomass", "p_nom": round(net_cap[area], 1)})
        report.append((area, cap_area[area], big_cap_area[area], net_cap[area], net_e[area] / 1e6, cf_mean))

    # --- Excel に書き込む（既存の中小バイオマスは置き換え） ---
    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb["generators"]
    hdr = {c.value: c.column for c in ws[1]}
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, hdr["name"]).value).startswith("Biomass (small"):
            ws.delete_rows(r)
    template = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, hdr["carrier"]).value == "Biomass")
    for row in new_rows:
        r = ws.max_row + 1
        for c in range(1, ws.max_column + 1):                              # 既存の大型バイオマスの行を雛形にする
            ws.cell(r, c).value = ws.cell(template, c).value
        for k, v in row.items():
            ws.cell(r, hdr[k]).value = v
    wc = wb["biomass_cf"]
    keep = [c for c in range(1, wc.max_column + 1) if not str(wc.cell(1, c).value).startswith("Biomass (small")]
    old = [[wc.cell(r, c).value for c in keep] for r in range(1, wc.max_row + 1)]
    idx = wb.sheetnames.index("biomass_cf")
    del wb["biomass_cf"]
    wc = wb.create_sheet("biomass_cf", index=idx)
    names = list(new_cols)
    for r, row in enumerate(old):
        wc.append(row + (names if r == 0 else [float(new_cols[n][r - 1]) for n in names]))
    wb.save(args.xlsx)

    print(f"{args.xlsx} に中小バイオマス {len(new_rows)} 件を追加")
    print(f"全国（沖縄を除く）: 容量 {cap_area.sum():,.0f} MW, {args.year}年の買取電力量 {energy_year/1e6:.2f} TWh")
    print(f"{'エリア':8s} {'FIT容量':>8s} {'大型分':>7s} {'追加容量':>8s} {'追加TWh':>7s} {'利用率':>6s}")
    for a, c, b, nc, e, cf in report:
        print(f"{a:10s} {c:8.0f} {b:7.0f} {nc:8.0f} {e:7.2f} {cf:6.1%}")
    print(f"合計追加: {sum(r[3] for r in report):,.0f} MW, {sum(r[4] for r in report):.2f} TWh")


if __name__ == "__main__":
    main()
