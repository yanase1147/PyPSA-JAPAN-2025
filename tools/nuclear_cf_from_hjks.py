"""HJKS のユニット別・1時間ごとの原子力出力CSVから、原子力発電機(発電所単位)のCFを作り、
入力Excel の nuclear_cf シートに書き込む。

入力CSV (例: nuclear_2025_hourly_by_unit.csv):
    time, <エリア>|<発電所名>|<号機>, ...      # 値は出力 [kW], 8760行
出力: Excel の nuclear_cf シート (8760行 x 原子力発電機数, 列名 = generators シートの発電機名)
    CF = 発電所内ユニット出力の合計 / generators シートの p_nom
    build_network はこのシートがあると、原子力の出力をCFどおりに固定する (p_min_pu = p_max_pu)。

使い方:
    python tools/nuclear_cf_from_hjks.py data/Network_config_2025.xlsx nuclear_2025_hourly_by_unit.csv
"""
import sys

import openpyxl
import pandas as pd

# HJKS の発電所名に含まれる語 -> generators シートの発電機名
PLANT = {
    "泊": "Tomari NPS", "東通": "Higashidori NPS", "女川": "Onagawa NPS",
    "東海第二": "Tokai No.2 NPS", "柏崎刈羽": "Kashiwazaki-Kariwa NPS", "浜岡": "Hamaoka NPS",
    "敦賀": "Tsuruga NPS Unit 2", "志賀": "Hokuriku EPCO Shika NPS", "美浜": "Mihama NPS",
    "高浜": "Takahama NPS", "大飯": "Ohi NPS", "島根": "Shimane NPS",
    "伊方": "Shikoku EPCO Ikata NPS", "川内": "Sendai NPS (Kagoshima)", "玄海": "Genkai NPS",
}
SHEET = "nuclear_cf"


def plant_of(col):
    hits = {v for k, v in PLANT.items() if k in col.split("|")[1]}
    if len(hits) != 1:
        raise ValueError(f"発電所を特定できない列: {col}")
    return hits.pop()


def main(xlsx, src):
    g = pd.read_excel(xlsx, sheet_name="generators")
    nu = g[g.carrier == "Nuclear"].set_index("name").p_nom

    d = pd.read_csv(src, encoding="utf-8-sig", parse_dates=["time"], index_col="time")
    if len(d) != 8760 or not (d.index.to_series().diff().dropna() == pd.Timedelta("1h")).all():
        raise ValueError(f"{src}: 1時間刻み8760行ではありません ({len(d)}行)")
    if d.isna().any().any() or (d < 0).any().any():
        raise ValueError(f"{src}: 欠損または負の値があります")

    mw = (d / 1e3).T.groupby(plant_of).sum().T          # kW -> MW, 発電所単位に集計
    missing = set(nu.index) - set(mw.columns)
    if missing:
        raise ValueError(f"CSVに無い原子力発電機: {sorted(missing)}")
    cf = (mw[nu.index] / nu).reset_index(drop=True)
    over = cf.max()[cf.max() > 1 + 1e-6]
    if len(over):
        print(f"警告: CF > 1 を 1 に丸めます {over.round(3).to_dict()}", file=sys.stderr)
    cf = cf.clip(0, 1).round(6)

    wb = openpyxl.load_workbook(xlsx)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, index=wb.sheetnames.index("biomass_cf") + 1)
    ws.append(list(cf.columns))
    for row in cf.itertuples(index=False):
        ws.append(list(row))
    wb.save(xlsx)

    e = cf.mul(nu).sum() / 1e6
    print(f"{xlsx} の {SHEET} シートを更新 ({cf.shape[0]}行 x {cf.shape[1]}列), 合計 {e.sum():,.1f} TWh")
    print(pd.DataFrame({"年平均CF": cf.mean().round(3), "TWh": e.round(2)}).to_string())


if __name__ == "__main__":
    main(*sys.argv[1:3])
