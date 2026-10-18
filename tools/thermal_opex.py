"""石炭・LNG火力の限界費用（燃料費）を、燃料価格と発電方式ごとの効率から計算するための
シートを入力Excel に書き込む。

使い方:
    python tools/thermal_opex.py data/Network_config_2025.xlsx

書き込むシート:
- thermal_efficiency: 発電機ごとの方式（容量の内訳）、効率、燃料。値は手で直してよい
- fuel_price: 燃料ごとの標準発熱量、石油石炭税、2025年の月別の輸入価格（円/t）
build_network は両シートがあると、限界費用を月別に
    限界費用 [円/MWh] = (輸入価格 + 税) [円/t] ÷ 標準発熱量 [GJ/t] × 3.6 [GJ/MWh] ÷ 効率 + 燃料費以外の変動費
として設定する（generators シートの marginal_cost は年平均値に書き換えて参考に残す）。

出典:
- 輸入価格: 財務省貿易統計（通関CIF）2025年1〜12月。LNG・一般炭（新電力ネット「通関統計」の集計表）
- 標準発熱量: 資源エネルギー庁「エネルギー源別標準発熱量」2023年度改訂案（輸入一般炭 25.88、LNG 54.69 GJ/t、総発熱量）
- 石油石炭税（地球温暖化対策税を含む）: LNG 1,860 円/t、石炭 1,370 円/t
- 効率（発電端・高位発熱量）: 方式ごとの目安。主な根拠は
    * 発電コスト検証WG（2021/4/12）: USC石炭 43.5%、1500℃級GTCC 54.5%（新設のモデルプラント）
    * 省エネ法 火力のベンチマーク（2018/12/3）: 新設基準 石炭 42.0%・LNG 50.5%、事業者目標 石炭 41%・LNG 48%
  で、既設の古い設備ほど低い値とした。石炭の方式（亜臨界・SC・USC・IGCC・PFBC）は経産省 石炭火力検討WG
  第1回 参考資料1（2020/8/7）の号機別一覧に合わせた。LNG の方式の内訳は公開情報にもとづく推定。
- 製鉄所の副生ガス（高炉ガス等）を主燃料とする発電所は、燃料を「副生ガス」とし、石炭と同じ熱量単価とした（仮定）。
"""
import sys

import openpyxl

# 方式ごとの効率（発電端・HHV）と燃料
CLASS = {
    # LNG
    "lng_steam":  (0.39, "lng", "LNG 従来型汽力"),
    "cc1100":     (0.44, "lng", "GTCC 1100℃級"),
    "cc1300":     (0.48, "lng", "GTCC 1300℃級"),
    "cc1500":     (0.52, "lng", "GTCC 1500℃級"),
    "cc1600":     (0.555, "lng", "GTCC 1600℃級"),
    "cc1650":     (0.57, "lng", "GTCC 1650℃級（J形）"),
    "lng_small":  (0.42, "lng", "小型GT・コージェネ等"),
    "byproduct":  (0.38, "byproduct", "製鉄所の副生ガス（汽力）"),
    # 石炭
    "sub":        (0.37, "coal", "亜臨界・小型"),
    "sc":         (0.395, "coal", "超臨界（SC）"),
    "usc":        (0.42, "coal", "超々臨界（USC）"),
    "usc_new":    (0.435, "coal", "USC（2015年以降）"),
    "igcc":       (0.46, "coal", "IGCC（商用）"),
    "pfbc":       (0.40, "coal", "PFBC"),
    "igcc_demo":  (0.40, "coal", "IGCC 実証機"),
}

# 発電機（行）ごとの方式の内訳: [(方式, 容量の割合), ...]。1つなら割合 1
PLANT = {
    # ---- LNG ----
    "Kawagoe Thermal PS": [("lng_steam", 1400 / 4802), ("cc1300", 1650 / 4802), ("cc1500", 1752 / 4802)],
    "Shin-Nagoya Thermal PS": [("cc1300", 1458 / 3058), ("cc1500", 1600 / 3058)],
    "Nishi-Nagoya Thermal PS": [("cc1600", 1)],
    "Chita No.2 Thermal PS": [("lng_steam", 1)],
    "Chita Thermal PS": [("lng_steam", 1)],
    "Yokkaichi Thermal PS": [("cc1100", 1)],
    "Fukuyama Kyodo PS": [("byproduct", 1)],
    "Yanai PS Group 2": [("cc1300", 1)],
    "Yanai PS Group 1 (Upper)": [("cc1100", 1)],
    "Yanai PS Group 1 (Lower)": [("cc1100", 1)],
    "Kurashiki Kyodo PS": [("byproduct", 1)],
    "Tamashima PS Unit 1": [("lng_steam", 1)],
    "Mizushima PS Unit 3": [("lng_steam", 1)],
    "Mizushima PS Unit 1": [("cc1300", 1)],
    "Ishikariwan Shinko PS": [("cc1600", 1)],
    "Toyama Shinko Thermal PS": [("cc1600", 0.45), ("lng_steam", 0.55)],
    "Himeji No.2 PS": [("cc1600", 1)],
    "Sakaiko PS": [("cc1500", 1)],
    "Himeji No.1 PS": [("cc1100", 1)],
    "Mitsubishi Takasago Works Demonstration Combined Cycle PS Unit 2": [("cc1650", 1)],
    "Senboku Natural Gas No.1 PS": [("cc1500", 1)],
    "Senboku Natural Gas No.2 PS": [("cc1500", 1)],
    "Funamachi PS": [("cc1300", 1)],
    "Torishima Energy Center PS": [("lng_small", 1)],
    "Shin-Oita PS": [("cc1100", 0.35), ("cc1300", 0.45), ("cc1600", 0.20)],
    "Shin-Kokura PS": [("lng_steam", 1)],
    "Tobata Kyodo PS": [("byproduct", 1)],
    "Oita Kyodo PS": [("byproduct", 1)],
    "Oita Steel Works PS Unit 9": [("byproduct", 1)],
    "Joetsu Thermal PS": [("cc1500", 1)],
    "Tohoku EPCO Higashi-Niigata Thermal PS Group 4": [("cc1500", 1)],
    "Tohoku EPCO Higashi-Niigata Thermal PS Group 3": [("cc1100", 1)],
    "Tohoku EPCO Higashi-Niigata Thermal PS Unit 2": [("lng_steam", 1)],    # 重油・原油・LNG の汽力（入力では石炭になっていた）
    "Fukushima Natural Gas PS": [("cc1600", 1)],
    "Tohoku EPCO Shin-Sendai Thermal PS Group 3": [("cc1600", 1)],
    "Tohoku EPCO Joetsu Thermal PS Unit 1": [("cc1650", 1)],
    "Tohoku EPCO Sendai Thermal PS Unit 4": [("cc1500", 1)],
    "Tohoku EPCO Hachinohe Thermal PS Unit 5": [("cc1300", 1)],
    "Futtsu Thermal PS": [("cc1100", 0.37), ("cc1500", 0.63)],
    "Chiba Thermal PS": [("cc1500", 0.66), ("cc1600", 0.34)],
    "Sodegaura Thermal PS": [("lng_steam", 1)],
    "Kawasaki Thermal PS": [("cc1500", 0.44), ("cc1600", 0.56)],
    "Anegasaki Thermal PS": [("cc1650", 0.62), ("lng_steam", 0.38)],
    "Yokohama Thermal PS": [("cc1300", 0.8), ("lng_steam", 0.2)],
    "Goi Thermal PS": [("cc1650", 1)],
    "Higashi-Ogishima Thermal PS": [("lng_steam", 1)],
    "Moka PS": [("cc1600", 1)],
    "Ogishima Power": [("cc1500", 1)],
    "Kimitsu Kyodo Thermal PS": [("byproduct", 1)],
    "Minami-Yokohama Thermal PS": [("lng_steam", 1)],
    "Shinagawa Thermal PS": [("cc1300", 1)],
    "Kawasaki Natural Gas PS": [("cc1500", 1)],
    "Kawasaki PS": [("cc1300", 1)],
    "JFE Steel Kawatetsu Chiba Clean Power Station": [("byproduct", 1)],
    "JFE Steel East Japan Works (Chiba) West PS": [("byproduct", 1)],
    "Yokosuka Power Station": [("lng_small", 1)],
    "JFE Steel East Japan Works (Keihin) Ogishima Thermal PS": [("byproduct", 1)],
    "JFE Steel East Japan Works (Chiba) Combined Cycle PS": [("byproduct", 1)],
    "Shin-Nakasode PS": [("cc1300", 1)],
    "Nakasode Clean Power PS": [("cc1300", 1)],
    "ZeroWatt Power Mihama PS": [("lng_small", 1)],
    "ZeroWatt Power Ichihara PS": [("lng_small", 1)],
    # ---- 石炭 ----
    "Hekinan Thermal PS": [("sc", 1400 / 4100), ("usc", 2700 / 4100)],      # 1・2号 SC、3〜5号 USC
    "Kobe PS": [("sc", 1400 / 2700), ("usc_new", 1300 / 2700)],             # 1・2号 SC、3・4号 USC（2022〜23年）
    "Nakoso PS": [("sub", 250 / 1450), ("sc", 1200 / 1450)],               # 7号 亜臨界、8・9号 SC
    "J-POWER Tachibanawan Thermal PS": [("usc", 1)],
    "Matsuura Thermal PS": [("sc", 0.5), ("usc", 0.5)],
    "Hitachinaka Thermal PS": [("usc", 1)],
    "Soma Kyodo Power Shinchi Thermal PS": [("sc", 1)],
    "Kashima Thermal PS": [("cc1500", 1)],                                  # JERA 鹿島 7号系列（都市ガス GTCC、方式は推定）
    "Maizuru PS": [("usc", 1)],
    "Matsuura PS": [("sc", 700 / 1700), ("usc_new", 1000 / 1700)],
    "Tomato-Atsuma PS": [("sub", 350 / 1650), ("sc", 600 / 1650), ("usc", 700 / 1650)],
    "Reihoku PS": [("sc", 0.5), ("usc", 0.5)],
    "Yokosuka Thermal PS": [("usc_new", 1)],
    "Hokuriku EPCO Tsuruga Thermal PS": [("sc", 500 / 1200), ("usc", 700 / 1200)],
    "Hokuriku EPCO Nanao-Ota Thermal PS": [("usc", 1)],                     # 1・2号とも USC
    "Taketoyo Thermal PS": [("usc_new", 1)],
    "Tohoku EPCO Haramachi Thermal PS Unit 1": [("usc", 1)],
    "Tohoku EPCO Haramachi Thermal PS Unit 2": [("usc", 1)],
    "Misumi PS Unit 1": [("usc", 1)],
    "Misumi PS Unit 2": [("usc_new", 1)],
    "Shikoku EPCO Saijo PS": [("usc_new", 500 / 750), ("sub", 250 / 750)],
    "Shikoku EPCO Tachibanawan PS": [("usc", 1)],
    "Sakata Kyodo Thermal PS": [("sub", 1)],
    "Hitachinaka Kyodo Thermal PS": [("usc_new", 1)],
    "Tohoku EPCO Noshiro Thermal PS Unit 1": [("sc", 1)],
    "Tohoku EPCO Noshiro Thermal PS Unit 2": [("usc", 1)],
    "Tohoku EPCO Noshiro Thermal PS Unit 3": [("usc_new", 1)],
    "J-POWER Takehara Thermal PS New Unit 1": [("usc_new", 1)],
    "J-POWER Isogo Thermal PS New Unit 2": [("usc", 1)],
    "Isogo Thermal PS New Unit 1": [("usc", 1)],                            # 石炭（入力では石油になっていた）
    "Takehara Thermal PS Unit 3": [("sc", 1)],                             # 石炭（入力では石油になっていた）
    "Hirono IGCC PS": [("igcc", 1)],
    "Nakoso IGCC PS": [("igcc", 1)],
    "Shin-Onoda PS Unit 1": [("sc", 1)],
    "Shin-Onoda PS Unit 2": [("sc", 1)],
    "Takasago Thermal PS": [("sub", 1)],
    "Matsushima Thermal PS": [("sc", 1)],
    "Nippon Steel Kashima PS": [("sc", 1)],                                 # 日本製鉄 鹿島 1号 SC（2007年）
    "Kashima Power Kashima Thermal PS Unit 2": [("usc_new", 1)],            # 鹿島パワー 2号 USC（2020年）
    "Hirono Thermal PS Units 5-6": [("usc", 1)],                            # JERA 広野 5号（2004年）・6号（2013年）USC
    "Osaki PS Unit 1-1": [("pfbc", 1)],
    "Oxygen-blown IGCC Demonstration PS": [("igcc_demo", 1)],
    "Kanda PS": [("pfbc", 1)],
}
# 上にない石炭の行（300 MW 未満の小型・自家発系）は "sub"

FUEL = {  # 標準発熱量 [GJ/t], 石油石炭税 [円/t], 2025年1〜12月の輸入価格 [円/t]
    "lng": (54.69, 1860, [100350, 94310, 90910, 88780, 86660, 85470, 85050, 84650, 82430, 81700, 84140, 85780]),
    "coal": (25.88, 1370, [23390, 21690, 19580, 18070, 17270, 16980, 16910, 17850, 17860, 18400, 17910, 18840]),
}
FUEL["byproduct"] = FUEL["coal"]

# 燃料費以外の変動費 [円/MWh]（価格水準の補正）。JEPX エリアプライス（2025年10〜12月）に合わせた較正値で、
# 運転維持費の実測値ではない。燃料費だけの限界費用では、西日本の価格が石炭の限界費用で決まり、JEPX より
# 2.4〜3.6円/kWh 低くなる。実際には、予備力の確保・最低出力・燃料制約・入札時の上乗せなどで価格が上がる分を、
# 石炭に 3.0、LNG に 1.0 円/kWh として一括で表す（LNGの最低出力 20% を入れたうえで、10〜12月の平均絶対誤差が最小の組み合わせ）。0 にすると燃料費だけの限界費用になる。
# run_dispatch.py の --vom-coal / --vom-lng で実行時に上書きできる。
VOM = {"lng": 1000, "coal": 3000, "byproduct": 3000}


def main(xlsx):
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["generators"]
    hdr = {c.value: c.column for c in ws[1]}
    rows = []
    for r in range(2, ws.max_row + 1):
        name, carrier = ws.cell(r, hdr["name"]).value, ws.cell(r, hdr["carrier"]).value
        if carrier not in ("Coal", "Gas"):
            continue
        mix = PLANT.get(name, [("sub", 1)] if carrier == "Coal" else None)
        if mix is None:
            raise ValueError(f"LNG の発電機 {name} の方式が未設定です")
        if abs(sum(s for _, s in mix) - 1) > 1e-6:
            raise ValueError(f"{name}: 方式の割合の合計が 1 ではありません")
        fuels = {CLASS[c][1] for c, _ in mix}
        if len(fuels) != 1:
            raise ValueError(f"{name}: 燃料の違う方式が混在しています")
        eff = sum(CLASS[c][0] * s for c, s in mix)
        desc = " + ".join(f"{CLASS[c][2]} {s:.0%}" for c, s in mix)
        rows.append((name, carrier, round(eff, 4), fuels.pop(), desc, r))

    # marginal_cost 列を年平均の値に更新（参考値。計算は fuel_price の月別値を使う）
    for name, carrier, eff, fuel, desc, r in rows:
        gj, tax, price = FUEL[fuel]
        mc = (sum(price) / 12 + tax) / gj * 3.6 / eff + VOM[fuel]
        ws.cell(r, hdr["marginal_cost"]).value = round(mc)
        ws.cell(r, hdr["efficiency"]).value = eff

    for sheet in ("thermal_efficiency", "fuel_price"):
        if sheet in wb.sheetnames:
            del wb[sheet]
    after = wb.sheetnames.index("interconnections")
    te = wb.create_sheet("thermal_efficiency", index=after)
    te.append(["name", "carrier", "efficiency", "fuel", "basis"])
    for name, carrier, eff, fuel, desc, _ in rows:
        te.append([name, carrier, eff, fuel, desc])
    fp = wb.create_sheet("fuel_price", index=after + 1)
    fp.append(["fuel", "heat_GJ_per_t", "tax_yen_per_t"] + [f"m{m:02d}" for m in range(1, 13)] + ["vom_yen_per_MWh"])
    for fuel, (gj, tax, price) in FUEL.items():
        fp.append([fuel, gj, tax] + price + [VOM[fuel]])
    wb.save(xlsx)

    # 確認用の集計
    import pandas as pd
    df = pd.DataFrame(rows, columns=["name", "carrier", "eff", "fuel", "desc", "r"])
    g = pd.read_excel(xlsx, sheet_name="generators").set_index("name")
    df["p_nom"] = df.name.map(g.p_nom)
    df["mc"] = df.name.map(g.marginal_cost)
    for fuel, d in df.groupby("fuel"):
        w = d.p_nom / d.p_nom.sum()
        print(f"{fuel:9s} {len(d):3d}行 {d.p_nom.sum()/1e3:5.1f} GW  効率(容量加重) {(d.eff*w).sum():.1%}  "
              f"限界費用(年平均) {d.mc.min()/1e3:.2f}〜{d.mc.max()/1e3:.2f} 円/kWh, 容量加重 {(d.mc*w).sum()/1e3:.2f}")


if __name__ == "__main__":
    main(sys.argv[1])
