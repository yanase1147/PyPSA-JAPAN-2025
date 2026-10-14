"""石炭火力の稼働率（利用可能率）から時間別の出力上限 p_max_pu を作り、
入力Excel の coal_cf シートに書き込む。

使い方:
    python tools/coal_availability.py data/Network_config_2025.xlsx [--year 2025]

考え方:
- 計画停止（定期点検）: 九州電力「料金算定の前提となる供給力について」(2012) による
    * ボイラー点検は2年に1回、大型石炭の停止は約60日 → 対象年はユニットの半数が60日停止
    * 夏前の簡易中間点検 約10日 → 春に60日点検を行うユニット以外は全ユニットで実施
    * 点検は需要の少ない春（3〜6月）・秋（9月中旬〜12月中旬）に置き、
      同じエリアのユニットの開始日を窓の中で均等にずらす
- 計画外停止: 広域機関 第115回 調整力及び需給バランス評価等に関する委員会 資料1-2 (2026/1)
    * 石炭の計画外停止率 6.9%（2022〜2024年度平均）を期待値として全時間の上限に掛ける
- generators シートには発電所単位にまとめた行（例: 碧南 4,100 MW）があるため、
  1行を round(p_nom / 800 MW) 基（最低1基）のユニットに分けて停止を割り当てる。
  1基が停止すると、その行の上限は 1/基数 だけ下がる。
- 停止の開始・終了で上限が急変するとランプ制約と矛盾するため、
  1時間あたりの上限の変化を MAX_STEP 以下に抑える（前後1〜2時間でなだらかにする）。
"""
import argparse

import numpy as np
import openpyxl
import pandas as pd

SHEET = "coal_cf"
UNIT_MW = 800            # 発電所単位の行を何基とみなすかの目安 [MW/基]
MAJOR_DAYS = 60          # 定期点検（2年に1回）の停止日数
MID_DAYS = 10            # 夏前の中間点検の停止日数
FORCED_OUTAGE = 0.069    # 計画外停止率
MAX_STEP = 0.45          # 1時間あたりの上限の変化の最大値（石炭の最小ランプ率0.5より小さく）

# 定期点検の開始日の範囲（この範囲で開始すれば窓内に収まる）
SPRING = ("03-01", "05-01")   # 3/1〜6/30
AUTUMN = ("09-16", "10-16")   # 9/16〜12/15
MID = ("05-15", "06-20")      # 夏前の中間点検 5/15〜6/30


def spread(start, end, k, year):
    """[start, end] の開始日を k 個、均等に並べる。"""
    s, e = pd.Timestamp(f"{year}-{start}"), pd.Timestamp(f"{year}-{end}")
    if k == 1:
        return [s + (e - s) / 2]
    return [s + (e - s) * i / (k - 1) for i in range(k)]


def build_units(coal):
    """発電所単位の行をユニットに分ける。"""
    rows = []
    for name, r in coal.iterrows():
        n = max(1, int(round(r.p_nom / UNIT_MW)))
        for j in range(n):
            rows.append({"gen": name, "area": r.area, "share": 1.0 / n, "mw": r.p_nom / n})
    u = pd.DataFrame(rows)
    # エリアごとに容量の大きい順に並べ、交互に「今年60日点検」(A) / 「中間点検のみ」(B)
    u = u.sort_values(["area", "mw"], ascending=[True, False]).reset_index(drop=True)
    u["rank"] = u.groupby("area").cumcount()
    u["major"] = u["rank"] % 2 == 0
    return u


def schedule(u, year):
    """各ユニットの停止期間 [(開始, 日数, 種別), ...] を決める。"""
    out = {i: [] for i in u.index}
    for area, grp in u.groupby("area"):
        major = grp[grp.major]
        # 60日点検: 春と秋に交互に振り分け、それぞれの窓の中で開始日を均等にずらす
        spring_idx, autumn_idx = list(major.index[0::2]), list(major.index[1::2])
        for idx, (a, b) in ((spring_idx, SPRING), (autumn_idx, AUTUMN)):
            for i, d in zip(idx, spread(a, b, len(idx), year)):
                out[i].append((d, MAJOR_DAYS, "major"))
        # 中間点検: 春に60日点検を行うユニット以外
        mid_idx = [i for i in grp.index if i not in spring_idx]
        for i, d in zip(mid_idx, spread(*MID, len(mid_idx), year)):
            out[i].append((d, MID_DAYS, "mid"))
    return out


def smooth(a, step):
    """1時間あたりの変化が step 以下になるよう、上限を下げる方向にだけ補正する。"""
    a = a.copy()
    for t in range(1, len(a)):                      # 上がるときは前の値 + step まで
        a[t] = min(a[t], a[t - 1] + step)
    for t in range(len(a) - 2, -1, -1):             # 下がるときは次の値 + step まで
        a[t] = min(a[t], a[t + 1] + step)
    return a


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--year", type=int, default=2025)
    args = ap.parse_args()

    g = pd.read_excel(args.xlsx, sheet_name="generators").set_index("name")
    coal = g[g.carrier == "Coal"]
    snaps = pd.date_range(f"{args.year}-01-01", periods=8760, freq="h")

    u = build_units(coal)
    sched = schedule(u, args.year)

    # ユニットの稼働状態 (1 = 運転可能) → 行ごとに容量で合成
    avail = pd.DataFrame(0.0, index=snaps, columns=coal.index)
    planned_days = []
    for i, r in u.iterrows():
        on = np.ones(len(snaps))
        days = 0
        for start, d, _ in sched[i]:
            m = (snaps >= start) & (snaps < start + pd.Timedelta(days=d))
            on[m] = 0
            days += d
        planned_days.append(days)
        avail[r.gen] += r.share * on
    u["planned_days"] = planned_days

    cf = avail * (1 - FORCED_OUTAGE)
    cf = cf.apply(lambda s: pd.Series(smooth(s.values, MAX_STEP), index=s.index))
    cf = cf.clip(0, 1).round(6).reset_index(drop=True)

    # Excel に書き込む（他のシートはそのまま）
    wb = openpyxl.load_workbook(args.xlsx)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    after = "nuclear_cf" if "nuclear_cf" in wb.sheetnames else "biomass_cf"
    ws = wb.create_sheet(SHEET, index=wb.sheetnames.index(after) + 1)
    ws.append(list(cf.columns))
    for row in cf.itertuples(index=False):
        ws.append(list(row))
    wb.save(args.xlsx)

    # 確認用の集計
    w = coal.p_nom
    fleet = (cf * w).sum(axis=1) / w.sum()
    planned = (u.planned_days * u.mw).sum() / u.mw.sum() / 365
    print(f"{args.xlsx} の {SHEET} シートを更新 ({cf.shape[0]}行 x {cf.shape[1]}列)")
    print(f"ユニット数 {len(u)}（発電所単位の行を分割）, うち今年60日点検 {u.major.sum()} 基")
    print(f"計画停止率 {planned:.1%}, 計画外停止率 {FORCED_OUTAGE:.1%} → 年平均の稼働率 {fleet.mean():.1%}")
    m = fleet.groupby(snaps.month).mean()
    print("月別の稼働率:", " ".join(f"{k}月 {v:.0%}" for k, v in m.items()))
    print("エリア別の最小稼働率:", ((cf * w).T.groupby(coal.area).sum().T / coal.groupby("area").p_nom.sum())
          .min().round(2).to_dict())


if __name__ == "__main__":
    main()
