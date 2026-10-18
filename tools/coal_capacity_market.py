"""非効率石炭火力の稼働抑制（容量市場の設備利用率50%）を、generators シートの max_annual_cf 列に書き込む。

使い方:
    python tools/coal_capacity_market.py data/Network_config_2025.xlsx [--threshold 0.42] [--cap 0.5]

考え方:
- 容量市場では、設計効率42%未満の石炭火力は、年間の設備利用率が50%を超えると容量確保契約金額が
  20%減額される（実需給2025年度〜。広域機関 容量市場の在り方等に関する検討会 第35回 資料3-2, 2021）。
  実需給2027年度からは「設備利用率50%以下」がリクワイアメントになる（広域機関, 2023/7/20）。
  減額を避けるため、事業者は年間50%以下に抑えて運転すると仮定し、年間の発電量の上限として扱う。
- 発電機（行）ごとの方式の内訳は tools/thermal_opex.py の PLANT を使う。効率が threshold 未満の方式の
  容量の割合を f として、行の上限を
      max_annual_cf = cap × f + 1.0 × (1 − f)
  とする（効率の高い部分は上限なし。点検・計画外停止は coal_cf シートで別に扱う）。
- 製鉄所の副生ガス発電（方式 byproduct）は石炭を主燃料としないので対象外。
- 工場の自家発（SELF_GENERATION）は対象外。発電の大半を自社の工場で使い、容量市場の落札電源ではないと
  考えられるため、容量確保契約金額の減額がかからない。IPP（卸供給）や小売向けに売電する小型の石炭は
  容量市場に参加していると考え、対象に含める。

build_network は max_annual_cf 列があると、発電機の年間発電量の上限 e_sum_max を
max_annual_cf × p_nom × 時間数 とする。ローリングホライズンでは、残りの上限を窓に配分して解く
（run_dispatch.py の rolling_horizon を参照）。
"""
import argparse
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from thermal_opex import CLASS, PLANT  # noqa: E402


# 工場の自家発（発電の大半を自社工場で消費）。経産省 石炭火力検討WG 第1回 参考資料1（2020）の事業者名で判断
SELF_GENERATION = {
    "Ube Industries PS",                # 宇部興産（現 UBE）宇部 5・6号
    "Tokuyama East PS No.3",            # トクヤマ 徳山製造所東
    "Sumitomo Osaka Cement Ako PS",     # 住友大阪セメント 赤穂工場
    "Nyugawa Thermal PS",               # 住友共同電力 壬生川（住友グループの工場に供給）
    "Niihama-Nishi Thermal PS",         # 住友共同電力 新居浜西
}


def annual_cap(name, carrier, threshold, cap):
    """行の max_annual_cf と、効率が threshold 未満の容量の割合を返す（対象外は None）。"""
    if carrier != "Coal" or name in SELF_GENERATION:
        return None, 0.0
    mix = PLANT.get(name, [("sub", 1)])
    if all(CLASS[k][1] != "coal" for k, _ in mix):          # 副生ガスなど
        return None, 0.0
    f = sum(w for k, w in mix if CLASS[k][1] == "coal" and CLASS[k][0] < threshold)
    if f <= 0:
        return None, 0.0
    return round(cap * f + (1 - f), 4), f


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--threshold", type=float, default=0.42, help="非効率とする設計効率の閾値（発電端・HHV）")
    ap.add_argument("--cap", type=float, default=0.5, help="非効率部分の年間設備利用率の上限")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb["generators"]
    hdr = {c.value: c.column for c in ws[1]}
    if "max_annual_cf" not in hdr:
        col = ws.max_column + 1
        ws.cell(1, col).value = "max_annual_cf"
        hdr["max_annual_cf"] = col
    n_cap, mw_low, mw_tot = 0, 0.0, 0.0
    for r in range(2, ws.max_row + 1):
        name, carrier = ws.cell(r, hdr["name"]).value, ws.cell(r, hdr["carrier"]).value
        p_nom = float(ws.cell(r, hdr["p_nom"]).value or 0)
        value, f = annual_cap(name, carrier, args.threshold, args.cap)
        ws.cell(r, hdr["max_annual_cf"]).value = value
        if carrier == "Coal":
            mw_tot += p_nom
        if value is not None:
            n_cap += 1
            mw_low += f * p_nom
    wb.save(args.xlsx)
    print(f"{args.xlsx}: 石炭 {n_cap} 行に max_annual_cf を設定"
          f"（効率{args.threshold:.0%}未満の容量 {mw_low/1e3:.1f} GW / 石炭 {mw_tot/1e3:.1f} GW、上限 {args.cap:.0%}）")


if __name__ == "__main__":
    main()
