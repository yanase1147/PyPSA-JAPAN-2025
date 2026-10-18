"""LNG・石炭火力の最低出力（運転を下げられる下限）を、generators シートの must_run_pu 列に書き込む。

使い方:
    python tools/lng_must_run.py data/Network_config_2025.xlsx [--pu 0.3] [--coal-pu 0.3]

考え方:
- 起動停止なしの経済負荷配分では、需要が低い時間にLNG火力が大きく下がり、実績よりLNGが少なく石炭が多くなる
  （LNG 4〜12月 152 TWh、実績 201 TWh。電力調査統計・事業用）。実際には、予備力の確保、エリア内の系統制約、
  燃料契約などでLNGを下げ切れないと考えられる。これらは起動停止（最低出力・起動費）では表現できず、
  線形化した起動停止ありと起動停止なしで石炭・LNGの発電量はほぼ同じだった（石炭 329.6/329.9、ガス 188.7/190.5 TWh）。
- 値は、実績の月別LNG発電量の最小（5月: 16.85 TWh = 平均 22.6 GW、LNG容量の約30%）を超えない範囲で、
  10%・20%・30% を比べて 30% とした（20%より石炭 −18 TWh、LNG +19 TWh）。時間ごとの実績の下限は未確認で、
  30% は実績の月別最小（5月）と同じ水準で、時間ごとの下限としては高めの仮定である。
- build_network は、起動停止を扱わない発電機（unit_commitment=False の火力を含む）の p_min_pu を
  max(p_min_pu, must_run_pu) にする。起動停止を扱う発電機は、起動停止の最低出力（p_min_pu）に従う。
- 石炭（--coal-pu、既定 0.3）: 軽負荷期に火力を止められない制約（下げ代の確保）を表す。入れると出力抑制が増える
  （0.39→1.69 TWh、実績は約2 TWh。2023年度）。点検・停止で出せない時間は、下限も出力上限まで下げる。
  全石炭が常に最低出力以上という強めの仮定で、軽負荷期に一部を止める運用は起動停止ありで確かめる必要がある。
- 0 にすると最低出力なしに戻る。run_dispatch.py の --lng-min-pu / --coal-min-pu で実行時に上書きできる。
"""
import argparse

import openpyxl


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--pu", type=float, default=0.3, help="LNG火力の最低出力[定格に対する割合]")
    ap.add_argument("--coal-pu", type=float, default=0.3, help="石炭火力の最低出力[定格に対する割合]")
    args = ap.parse_args()
    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb["generators"]
    hdr = {c.value: c.column for c in ws[1]}
    if "must_run_pu" not in hdr:
        hdr["must_run_pu"] = ws.max_column + 1
        ws.cell(1, hdr["must_run_pu"]).value = "must_run_pu"
    n, mw = 0, 0.0
    for r in range(2, ws.max_row + 1):
        car = ws.cell(r, hdr["carrier"]).value
        gas = car == "Gas"
        ws.cell(r, hdr["must_run_pu"]).value = args.pu if gas else (args.coal_pu if car == "Coal" else None)
        if gas:
            n += 1
            mw += float(ws.cell(r, hdr["p_nom"]).value or 0)
    wb.save(args.xlsx)
    print(f"{args.xlsx}: LNG {n} 行（{mw/1e3:.1f} GW）の must_run_pu = {args.pu}、石炭 {args.coal_pu}")


if __name__ == "__main__":
    main()
