"""火力発電機の起動停止パラメータ（最低出力・最小運転/停止時間・起動費）を
入力Excel の generators シートに書き込む。

使い方:
    python tools/thermal_uc_params.py data/Network_config_2025.xlsx

設定値と出典:
- 最低出力（運転中の出力下限、定格比）
    * 石炭 500 MW以上: 30%
        発電コスト検証WG 第4回 資料2「系統安定化費用について」(2021/4/20): 石炭・GTCCの最低負荷30%
        火力原子力発電技術協会 (2021/5/27): USC石炭の最低出力は現状約30%
    * 石炭 500 MW未満: 50%
        資源エネルギー庁「出力制御の低減に向けた取組について」(2021/6/8): 電源Ⅲの石炭42事業者のうち
        24事業者が最低出力31%以上（うち8事業者は51%以上）。系統連系技術要件の既設の上限は50%
    * LNG: 30%  発電コスト検証WG (2021) のGTCC 30%。最新鋭GTCCは約25%（グリッドコード検討会 第14回 資料5, 2023/10/6）
    * 石油・その他: 30%  電源Ⅲの石油は大半が30%以下まで抑制可能（資源エネルギー庁 2021/6/8）
- 起動費
    * 石炭: 15万円/万kW・回 = 15,000 円/MW・回   発電コスト検証WG (2021) 事業者ヒアリングに基づく前提
    * LNG: 5万円/万kW・回 = 5,000 円/MW・回     同上
    * 石油・その他: LNGと同じと仮定
- 最小運転時間・最小停止時間（公表値が見つからないため、運用実態に基づく仮定）
    * 石炭: 運転24時間 / 停止12時間（週単位の起動停止 WSS が中心で、日単位の起動停止は行わない想定）
    * LNG・石油・その他: 運転4時間 / 停止4時間（グリッドコードの日間起動停止は解列から並列まで8時間以内）
"""
import sys

import openpyxl

COAL_LARGE_MW = 500
PARAMS = {
    #          p_min_pu, min_up_time, min_down_time, start_up_cost [円/MW]
    "Coal_L": (0.30, 24, 12, 15000),
    "Coal_S": (0.50, 24, 12, 15000),
    "Gas":    (0.30, 4, 4, 5000),
    "Oil":    (0.30, 4, 4, 5000),
    "Other":  (0.30, 4, 4, 5000),
}
COLS = ["committable", "p_min_pu", "min_up_time", "min_down_time", "start_up_cost"]


def main(xlsx):
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["generators"]
    hdr = {c.value: c.column for c in ws[1]}
    for col in COLS:                        # 無い列は末尾に追加
        if col not in hdr:
            c = ws.max_column + 1
            ws.cell(1, c).value = col
            hdr[col] = c
    n = {}
    for r in range(2, ws.max_row + 1):
        carrier = ws.cell(r, hdr["carrier"]).value
        p_nom = ws.cell(r, hdr["p_nom"]).value
        key = ("Coal_L" if p_nom >= COAL_LARGE_MW else "Coal_S") if carrier == "Coal" else carrier
        if key in PARAMS:
            pmin, up, down, su = PARAMS[key]
            vals = {"committable": True, "p_min_pu": pmin, "min_up_time": up,
                    "min_down_time": down, "start_up_cost": su * p_nom}
        else:                               # 火力以外は起動停止を扱わない
            vals = {"committable": False, "min_up_time": 0, "min_down_time": 0, "start_up_cost": 0}
        for col, v in vals.items():
            ws.cell(r, hdr[col]).value = v
        n[key] = n.get(key, 0) + 1
    wb.save(xlsx)
    print(f"{xlsx} の generators シートを更新:", {k: v for k, v in n.items() if k in PARAMS})


if __name__ == "__main__":
    main(sys.argv[1])
